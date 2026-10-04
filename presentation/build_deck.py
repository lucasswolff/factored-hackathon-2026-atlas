"""Build the six-slide Atlas pitch as an editable PPTX and a PDF preview.

Run with a Python environment containing python-pptx, reportlab, and Pillow:
    python presentation/build_deck.py
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from reportlab.lib.colors import HexColor
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas


ROOT = Path(__file__).resolve().parent
PPTX = ROOT / "atlas_factored_hackathon_2026.pptx"
PDF = ROOT / "atlas_factored_hackathon_2026.pdf"
CHAT = ROOT / "assets" / "spanish_conversation_crop.png"
REPO = "https://github.com/lucasswolff/factored-hackathon-2026-atlas"
DEMO = "https://55p7p6snz2nmfwiust5i5k7kmi0mjbqc.lambda-url.us-east-2.on.aws/"

W, H = 13.333, 7.5
BG = "F7F9F6"
WHITE = "FFFFFF"
NAVY = "112B38"
TEAL = "176F70"
MUTED = "617780"
LINE = "D7E3E1"
LIME = "D5F377"
PALE = "E9F4EE"
PALE_BLUE = "E8F2F6"
CORAL = "FFB897"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


pdfmetrics.registerFont(TTFont("DejaVu", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"))
pdfmetrics.registerFont(TTFont("DejaVu-Bold", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"))


class Deck:
    def __init__(self) -> None:
        self.pptx = Presentation()
        self.pptx.slide_width = Inches(W)
        self.pptx.slide_height = Inches(H)
        self.pdf = canvas.Canvas(str(PDF), pagesize=(W * 72, H * 72))
        self.slide = None
        self.page = 0

    def start(self, *, dark: bool = False) -> None:
        self.slide = self.pptx.slides.add_slide(self.pptx.slide_layouts[6])
        self.page += 1
        self.rect(0, 0, W, H, NAVY if dark else BG)
        if not dark:
            self.rect(0, 0, W, 0.09, TEAL)
            self.text(0.56, 0.27, 2.0, 0.25, "✳  ATLAS", 13, bold=True, color=NAVY)
            self.text(9.75, 0.29, 3.0, 0.22, "FACTORED HACKATHON 2026", 9, bold=True, color=TEAL, align="right")
            self.line(0.56, 7.11, 12.78, 7.11, LINE, 1)
            self.text(0.56, 7.19, 8.0, 0.17, "ATLAS  /  CREDIT-PRODUCT INFO & ELIGIBILITY DEMO", 7.8, color=MUTED)
            self.text(12.1, 7.18, 0.68, 0.18, f"{self.page:02d} / 06", 8, bold=True, color=MUTED, align="right")

    def finish(self) -> None:
        self.pdf.showPage()

    def save(self) -> None:
        self.pptx.save(PPTX)
        self.pdf.save()

    def rect(self, x: float, y: float, w: float, h: float, fill: str, *, stroke: str | None = None, radius: float = 0) -> None:
        shape = self.slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
            Inches(x), Inches(y), Inches(w), Inches(h),
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
        shape.line.fill.background() if stroke is None else None
        if stroke:
            shape.line.color.rgb = rgb(stroke)
            shape.line.width = Pt(1)
        if radius:
            try:
                shape.adjustments[0] = 0.12
            except Exception:
                pass
        self.pdf.setFillColor(HexColor(f"#{fill}"))
        if stroke:
            self.pdf.setStrokeColor(HexColor(f"#{stroke}"))
        if radius:
            self.pdf.roundRect(x * 72, (H - y - h) * 72, w * 72, h * 72,
                               radius * 72, fill=1, stroke=int(stroke is not None))
        else:
            self.pdf.rect(x * 72, (H - y - h) * 72, w * 72, h * 72,
                          fill=1, stroke=int(stroke is not None))

    def line(self, x1: float, y1: float, x2: float, y2: float, color: str, width: float = 1) -> None:
        shape = self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                            Inches(x1), Inches(y1), Inches(max(0.01, x2 - x1)), Inches(max(0.01, y2 - y1)))
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(color)
        shape.line.fill.background()
        self.pdf.setStrokeColor(HexColor(f"#{color}"))
        self.pdf.setLineWidth(width)
        self.pdf.line(x1 * 72, (H - y1) * 72, x2 * 72, (H - y2) * 72)

    def text(self, x: float, y: float, w: float, h: float, value: str, size: float,
             *, bold: bool = False, color: str = NAVY, align: str = "left", link: str | None = None,
             leading: float = 1.27) -> None:
        box = self.slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = box.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = Inches(0)
        frame.margin_top = frame.margin_bottom = Inches(0)
        frame.vertical_anchor = MSO_ANCHOR.TOP
        for i, line in enumerate(value.split("\n")):
            para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            para.text = line
            para.space_after = Pt(0)
            para.line_spacing = 1.0
            para.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
            for run in para.runs:
                run.font.name = "Arial"
                run.font.size = Pt(size)
                run.font.bold = bold
                run.font.color.rgb = rgb(color)
                if link:
                    run.hyperlink.address = link

        self.pdf.setFillColor(HexColor(f"#{color}"))
        self.pdf.setFont("DejaVu-Bold" if bold else "DejaVu", size)
        for i, line in enumerate(value.split("\n")):
            baseline = (H - y) * 72 - size * (0.97 + i * leading)
            if align == "right":
                self.pdf.drawRightString((x + w) * 72, baseline, line)
            elif align == "center":
                self.pdf.drawCentredString((x + w / 2) * 72, baseline, line)
            else:
                self.pdf.drawString(x * 72, baseline, line)
        if link:
            self.pdf.linkURL(link, (x * 72, (H - y - h) * 72, (x + w) * 72, (H - y) * 72), relative=0)

    def picture(self, path: Path, x: float, y: float, w: float, h: float) -> None:
        self.slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))
        self.pdf.drawImage(str(path), x * 72, (H - y - h) * 72, w * 72, h * 72)

    def title(self, kicker: str, heading: str, subheading: str | None = None) -> None:
        self.text(0.56, 0.83, 8.9, 0.2, kicker, 10, bold=True, color=TEAL)
        self.text(0.56, 1.06, 12.15, 0.53, heading, 26, bold=True)
        if subheading:
            self.text(0.56, 1.58, 12.1, 0.32, subheading, 12, color=MUTED)


def label(deck: Deck, x: float, y: float, value: str, *, fill: str = LIME, width: float = 1.55) -> None:
    deck.rect(x, y, width, 0.3, fill, radius=0.11)
    deck.text(x + 0.09, y + 0.065, width - 0.18, 0.17, value, 8.5, bold=True, color=NAVY)


def numbered_card(deck: Deck, x: float, y: float, w: float, number: str, heading: str,
                  lines: str, *, fill: str = WHITE, h: float = 1.34) -> None:
    deck.rect(x, y, w, h, fill, stroke=LINE, radius=0.14)
    deck.rect(x + 0.17, y + 0.18, 0.4, 0.4, LIME, radius=0.19)
    deck.text(x + 0.17, y + 0.275, 0.4, 0.12, number, 11, bold=True, align="center")
    deck.text(x + 0.7, y + 0.2, w - 0.85, 0.26, heading, 14, bold=True)
    deck.text(x + 0.7, y + 0.55, w - 0.88, h - 0.64, lines, 10.3, color=MUTED)


def build() -> None:
    if not CHAT.exists():
        raise SystemExit(f"Missing real conversation capture: {CHAT}")
    d = Deck()

    # 1. Cover
    d.start(dark=True)
    d.rect(0.0, 0.0, 0.12, H, LIME)
    d.rect(9.35, -0.55, 5.1, 5.1, TEAL, radius=2.3)
    d.rect(10.15, 0.25, 3.5, 3.5, NAVY, stroke=LIME, radius=1.6)
    d.text(10.75, 1.3, 2.3, 1.1, "✳", 79, bold=True, color=LIME, align="center")
    label(d, 0.7, 0.68, "AI-FIRST BANKING", fill=LIME, width=2.16)
    d.text(0.7, 1.45, 7.4, 1.0, "ATLAS", 63, bold=True, color=WHITE)
    d.text(0.74, 2.52, 7.55, 1.22, "Understand a card. Find a fit.\nTake a verified next step.", 27, bold=True, color=WHITE)
    d.text(0.74, 4.17, 7.1, 0.62, "Bilingual card guidance, simulated eligibility,\nand a path to human review in one conversation.", 15, color="C6D8D5")
    d.line(0.74, 5.47, 7.65, 5.47, "638987", 1)
    d.text(0.74, 5.76, 6.8, 0.35, "FACTORED HACKATHON 2026", 15, bold=True, color=LIME)
    d.text(0.74, 6.18, 6.8, 0.35, "Atlas  ·  Lucas Wolff", 17, color=WHITE)
    d.text(0.74, 7.11, 7.2, 0.18, "Working prototype  ·  Team-created offers and fictional public fixtures", 8.5, color="AFC4C1")
    d.finish()

    # 2. Product journey, with a concise data boundary
    d.start()
    d.title("01  /  THE PRODUCT JOURNEY", "From a card question to a customer-controlled next step",
            "Spanish and Portuguese chat helps a visitor understand the offers before choosing whether to act.")
    d.rect(0.56, 2.12, 12.22, 0.7, NAVY, radius=0.15)
    d.text(0.84, 2.32, 11.64, 0.27,
           "Clear terms now  ·  relevant choices next  ·  verified status for every mock action",
           15, bold=True, color=WHITE)
    for x, number, head, body, fill in [
        (0.56, "1", "Ask", "Fees and benefits\nin natural language", WHITE),
        (3.06, "2", "Compare", "See options and\ntheir trade-offs", PALE_BLUE),
        (5.56, "3", "Choose", "Keep the right card\nin context", WHITE),
        (8.06, "4", "Consent", "Allow an optional\ncard-specific precheck", PALE),
        (10.56, "5", "Next step", "Confirm a mock request\nor ask for review", WHITE),
    ]:
        d.rect(x, 3.18, 2.22, 1.88, fill, stroke=LINE, radius=0.14)
        d.rect(x + 0.16, 3.35, 0.38, 0.38, LIME, radius=0.18)
        d.text(x + 0.16, 3.45, 0.38, 0.12, number, 10.5, bold=True, align="center")
        d.text(x + 0.16, 3.92, 1.9, 0.3, head, 16, bold=True)
        d.text(x + 0.16, 4.35, 1.9, 0.55, body, 10.3, color=MUTED)
    d.rect(0.56, 5.43, 12.22, 0.61, PALE, stroke=LINE, radius=0.12)
    d.text(0.78, 5.61, 11.8, 0.25,
           "Four choices: Campus for students  ·  Horizon for daily use  ·  Rewards for travel  ·  Summit for premium travel",
           11.2, bold=True, color=TEAL)
    d.text(0.56, 6.27, 12.12, 0.4,
           "Data → value: campaigns shape entry; customer/product rows shape local tests; agent records shape mock routing; transcripts guide tone.\nPublic fixtures, offers, policy, and actions are team-created. Historical clicks show engagement, not card sales.",
           9.0, color=MUTED)
    d.finish()

    # 3. Real Spanish conversation
    d.start()
    d.title("02  /  WORKING EXPERIENCE", "See a live card question answered in Spanish",
            "Captured from the deployed judge app on 3 October 2026 · Rewards campaign · México · fictional fixture P02")
    d.rect(0.55, 2.02, 9.6, 4.72, WHITE, stroke=LINE, radius=0.13)
    d.picture(CHAT, 0.64, 2.11, 9.42, 4.54)
    d.rect(10.38, 2.02, 2.4, 4.72, NAVY, radius=0.15)
    d.text(10.64, 2.28, 1.9, 0.52, "Customer\nvalue", 19, bold=True, color=WHITE)
    d.line(10.64, 3.12, 12.52, 3.12, "597A7A", 1)
    d.text(10.64, 3.4, 1.92, 0.75, "Selected card\nstays in context", 13, bold=True, color=LIME)
    d.text(10.64, 4.22, 1.92, 0.75, "Follow-up gets\nan answer", 13, bold=True, color=LIME)
    d.text(10.64, 5.04, 1.92, 0.75, "Source IDs make\nclaims traceable", 13, bold=True, color=LIME)
    d.text(10.64, 6.18, 1.92, 0.31, "Live fictional-fixture chat", 9.2, color="C6D8D5")
    d.finish()

    # 4. Architecture
    d.start()
    d.title("03  /  HOW IT WORKS", "How Atlas turns a question into a safe next step",
            "Each part supports the journey: understand intent, ground facts, protect consent, and verify actions.")
    for x, w, head, body, fill in [
        (0.56, 2.34, "Customer starts", "Campaign / direct entry\nFictional demo customer", PALE_BLUE),
        (3.23, 2.52, "Lambda service", "Keeps the active card\nChecks access + limits", WHITE),
        (6.08, 3.06, "Haiku intent", "Fact / compare / apply\nNo private score sent", PALE),
    ]:
        d.rect(x, 2.15, w, 1.22, fill, stroke=LINE, radius=0.14)
        d.text(x + 0.2, 2.36, w - 0.4, 0.28, head, 16, bold=True)
        d.text(x + 0.2, 2.72, w - 0.4, 0.52, body, 10.5, color=MUTED)
    d.text(2.92, 2.53, 0.25, 0.3, "→", 21, color=TEAL)
    d.text(5.77, 2.53, 0.25, 0.3, "→", 21, color=TEAL)
    d.rect(9.49, 2.15, 3.29, 1.22, NAVY, radius=0.14)
    d.text(9.7, 2.36, 2.88, 0.28, "Durable state", 16, bold=True, color=WHITE)
    d.text(9.7, 2.72, 2.85, 0.52, "DynamoDB sessions\nMock actions + read-back", 10.5, color="C9D9D8")
    d.text(9.19, 2.53, 0.25, 0.3, "→", 21, color=TEAL)

    for x, w, head, body, fill in [
        (0.56, 3.75, "Useful answers", "Sonnet + versioned public facts\nSource IDs trace product claims", WHITE),
        (4.48, 3.75, "Controlled suggestion", "Selected fixture → policy suggestion\nConsent → synthetic precheck\nNo automatic approval", PALE),
        (8.40, 4.38, "Verified next step", "Separate confirmation → write → read back\nMock application: PENDING_REVIEW\nFull-thread handoff; bot pauses", "FFF3E8"),
    ]:
        d.rect(x, 3.87, w, 1.61, fill, stroke=LINE, radius=0.14)
        d.text(x + 0.21, 4.1, w - 0.42, 0.27, head, 15, bold=True)
        d.text(x + 0.21, 4.48, w - 0.42, 0.74, body, 10.5, color=MUTED)

    d.rect(0.56, 5.88, 12.22, 0.79, NAVY, radius=0.13)
    d.text(0.82, 6.07, 11.72, 0.46,
           "Trust: private score stays in the service  ·  card-specific consent gates precheck  ·  read-back gates success claims",
           11.5, bold=True, color=WHITE)
    d.finish()

    # 5. Current checks and their limits
    d.start()
    d.title("04  /  MEASURED BEHAVIOR", "A working journey, backed by measured checks",
            "Current 4 October build · product and safety paths checked in production.")
    numbered_card(d, 0.56, 2.05, 5.84, "1", "Product facts + context",
                  "Benefits and monthly fee came together with citations.\n'This card' kept Horizon across turns. A comparison\nagainst Summit returned lower-cost options.", h=1.72)
    numbered_card(d, 6.61, 2.05, 6.17, "2", "Intent + consent",
                  "A new application request reached card-specific consent.\nAn unnamed other card prompted a card choice.\nPrecheck decline led to a separate application choice.", h=1.72)
    numbered_card(d, 0.56, 4.02, 5.84, "3", "Profile + service boundaries",
                  "Student suggestion named the selected demo profile and\nits unverified enrollment. Cancellation got a support\npath; a second application got the one-per-chat rule.", h=1.72)
    numbered_card(d, 6.61, 4.02, 6.17, "4", "Checks behind the demo",
                  "Current build: 87 local tests; 7/8* live smoke flows.\nEarlier rehearsal: 12 entry paths, 3 application\nread-backs, 24 HTTP requests from six visitors.", h=1.72)
    d.rect(0.56, 6.03, 12.22, 0.81, PALE, stroke=LINE, radius=0.11)
    d.text(0.78, 6.13, 11.75, 0.18,
           "* After Rewards and Summit, 'a cheaper one' (PT: 'um mais barato') did not trigger a question about the reference card.",
           9.2, bold=True, color=NAVY)
    d.text(0.78, 6.36, 11.75, 0.18,
           "It listed alternatives; a separate 'cheaper than Rewards' answer even included the more expensive Summit.",
           9.2, color=NAVY)
    d.text(0.78, 6.62, 11.75, 0.14,
           "Separate, noncumulative checks  ·  400 model attempts/day cap",
           8.2, bold=True, color=TEAL)
    d.finish()

    # 6. Roadmap to a real customer service
    d.start()
    d.title("05  /  PATH TO A REAL SERVICE", "From judge demo to a bank-ready advisor",
            "The same journey can serve real customers when its facts, identity, and human support are bank-operated.")
    numbered_card(d, 0.56, 2.04, 5.84, "1", "Accurate live offers",
                  "Governed card terms and credit policy replace\nteam-created demo offers and thresholds.", h=1.43)
    numbered_card(d, 6.61, 2.04, 6.17, "2", "Trusted customer context",
                  "Verified identity, permission, and current account\ndata make personalized suggestions usable.", h=1.43)
    numbered_card(d, 0.56, 3.7, 5.84, "3", "A human who can finish",
                  "Live agent queue, full conversation handoff,\nreply in the same chat, and service levels.", h=1.43)
    numbered_card(d, 6.61, 3.7, 6.17, "4", "Quality customers can rely on",
                  "Independent bilingual evaluation, security,\ncapacity and cost budgets, and monitoring.", h=1.43)
    d.rect(0.56, 5.45, 12.22, 1.22, NAVY, radius=0.14)
    d.text(0.85, 5.72, 11.62, 0.36, "Today: a public demo. Next: approved bank systems and live reviewers.",
           17, bold=True, color=WHITE)
    d.text(0.85, 6.19, 4.7, 0.25, "Open the live demo  ↗", 11, bold=True, color=LIME, link=DEMO)
    d.text(6.6, 6.19, 5.72, 0.25, "Open the public repository  ↗", 11, bold=True,
           color=LIME, link=REPO, align="right")
    d.finish()

    d.save()
    print(PPTX)
    print(PDF)


if __name__ == "__main__":
    build()
